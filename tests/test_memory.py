"""长期记忆（v0.21）：资料库的解析和入库、指令结论、kb_search 工具、规划时自动召回、编码 agent 的检索工具。"""
import json
import threading

import pytest

from opc import coding
from opc.agents import KB_TOOL, call_kb, system_prompt
from opc.company import CompanyStore
from opc.knowledge import service
from opc.knowledge.base import KnowledgeBase
from opc.knowledge.company import directive_documents, library_root
from opc.knowledge.loaders import LoadError, NoText, load_text
from opc.worker import run

from fakes import call, reply
from harness import Env


def make_pdf(pages: list[list[str]]) -> bytes:
    """手写一个最小的 PDF：每页若干行英文（内置 Helvetica 字体），pypdf 能提取出文字。"""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for lines in pages:
        ops = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({ln}) Tj T*" for ln in lines) + " ET"
        objs.append(f"<< /Length {len(ops)} >>\nstream\n{ops}\nendstream")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {len(objs)} 0 R "
                    "/Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = b"%PDF-1.4\n", []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def make_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("Brand Guide", level=1)
    d.add_paragraph("Logo must keep clear space.")
    d.add_heading("Colors", level=2)
    d.add_paragraph("Primary color is ocean blue.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Name", "Hex"
    t.cell(1, 0).text, t.cell(1, 1).text = "Ocean", "#1F6FB2"
    d.save(str(path))


@pytest.fixture
def store(tmp_path):
    return CompanyStore(tmp_path / "home")


@pytest.fixture
def company(store):
    return store.create("公司", "", [{"id": "pm"}])


# ---------------------------------------------------------------- 解析

def test_pdf_pages_headers_and_page_numbers(tmp_path):
    p = tmp_path / "guide.pdf"
    p.write_bytes(make_pdf([["ACME Confidential", f"Refund rule number {i} applies here", str(i)] for i in range(1, 5)]))
    text = load_text(p)
    assert "## 第 1 页" in text and "## 第 4 页" in text and "Refund rule number 3" in text
    assert "ACME Confidential" not in text           # 每页都有的页眉去掉
    assert "\n3\n" not in text + "\n"                 # 单独的页码行去掉


def test_scanned_pdf_and_broken_files(tmp_path):
    empty = tmp_path / "scan.pdf"
    empty.write_bytes(make_pdf([[], []]))
    with pytest.raises(NoText):
        load_text(empty)
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"not a pdf at all")
    with pytest.raises(LoadError):
        load_text(broken)
    old = tmp_path / "old.docx"
    old.write_bytes(b"\xd0\xcf\x11\xe0 old word")
    with pytest.raises(LoadError):
        load_text(old)


def test_docx_headings_and_tables(tmp_path):
    p = tmp_path / "brand.docx"
    make_docx(p)
    text = load_text(p)
    assert "# Brand Guide" in text and "## Colors" in text and "Ocean | #1F6FB2" in text


# ---------------------------------------------------------------- 资料库

def test_library_upload_list_search_and_remove(company, store, tmp_path):
    assert service.add_document(company, "company", "退款规则.md", "# 退款\n\n7 天内无理由退款，金额原路退回。".encode())["chunks"] == 1
    info = service.add_document(company, "personal", "../../evil/简历.pdf", make_pdf([["Senior engineer resume"]]))
    assert info["name"] == "简历.pdf" and info["chunks"] >= 1   # 路径去掉，只留文件名
    bad = service.add_document(company, "company", "扫描件.pdf", make_pdf([[]]))
    assert bad["chunks"] == 0 and "没有提取到文字" in bad["error"]
    with pytest.raises(ValueError):
        service.add_document(company, "company", "表格.xlsx", b"x")
    lib = service.list_library(company)
    assert [i["name"] for i in lib["company"]] == ["扫描件.pdf", "退款规则.md"] and lib["personal"][0]["name"] == "简历.pdf"
    hits = service.search(company, "退款金额怎么退", "library")
    assert hits and hits[0]["source"] == "lib:退款规则.md"
    assert service.label(hits[0]).startswith("公司资料《退款规则.md》第 1-")
    assert service.search(company, "退款金额怎么退", "files") == []   # 资料不在文件区里

    # 个人资料所有公司共用；公司资料只属于自己
    other = store.create("另一家", "", [{"id": "pm"}])
    lib2 = service.list_library(other)
    assert lib2["company"] == [] and [i["name"] for i in lib2["personal"]] == ["简历.pdf"]

    # 解析失败的原因重启后还在
    service.reset()
    again = {i["name"]: i for i in service.list_library(company)["company"]}
    assert "没有提取到文字" in again["扫描件.pdf"]["error"]
    assert service.remove_document(company, "company", "退款规则.md")
    assert service.search(company, "退款金额怎么退", "library") == []
    assert not (library_root(company, "company") / "退款规则.md").exists()


def test_pdf_label_cites_page(company):
    long = ["Shipping policy line %d keeps parcels moving on time" % i for i in range(5)]
    service.add_document(company, "company", "policy.pdf",
                         make_pdf([long, ["Refund within seven days of delivery"] + long[:4]]))
    r = service.search(company, "refund delivery", "library")[0]
    assert service.label(r) == "公司资料《policy.pdf》第 2 页"


def test_merged_small_chunks_keep_all_titles():
    """很短的几节合成一块时，标题把每一节都写上（引用出处不会指错）；短的几页合成"第 a-b 页"。"""
    from opc.knowledge.chunking import chunk_file, merge_titles
    text = "# 规范\n\n## 一、金额\n金额用分。\n\n## 二、优惠券\n最长 90 天。\n\n## 三、积分\n12 个月过期。"
    assert [c["title"] for c in chunk_file("x.md", text)] == ["规范 / 一、金额；二、优惠券；三、积分"]
    assert merge_titles("第 2 页", "第 3 页") == "第 2-3 页" and merge_titles("第 2-3 页", "第 4 页") == "第 2-4 页"
    assert merge_titles("标题", "标题 / 1. 金额") == "标题 / 1. 金额"


# ---------------------------------------------------------------- 指令结论

def test_directive_document_collects_decisions():
    tasks = {"d1-a-t1": {"id": "d1-a-t1", "directive_id": "d1", "title": "写 PRD", "assignee": "pm", "status": "done",
                         "headline": "金额统一用分存储", "decisions": ["是否支持外币"], "ceo_note": "先不做外币",
                         "files": ["docs/PRD.md"], "result": "完成"}}
    msgs = [{"id": "c1", "sender": "ceo", "kind": "text", "text": "做一个记账 App", "directive_id": "d1", "ts": "2026-10-19"},
            {"id": "p1", "sender": "planner", "kind": "plan", "text": "先写 PRD", "directive_id": "d1"},
            {"id": "i1", "sender": "ceo", "kind": "interject", "text": "金额一律用分", "directive_id": "d1"},
            {"id": "q1", "sender": "engineer", "kind": "text", "text": "金额精度？", "consult": True, "task_id": "d1-a-t1"},
            {"id": "a1", "sender": "pm", "kind": "text", "text": "精确到分", "consult": True, "reply_to": "q1",
             "task_id": "d1-a-t1"},
            {"id": "r1", "sender": "reporter", "kind": "report", "text": "全部完成", "directive_id": "d1"}]
    title, text = directive_documents({"tasks": tasks, "channels": {"all": msgs}})["directive:d1"]
    assert title == "指令《做一个记账 App》"
    for part in ("已完成", "## CEO 中途的插话和补充", "金额一律用分", "## 计划", "精确到分", "金额统一用分存储",
                 "需要 CEO 决定：是否支持外币", "CEO 裁决：先不做外币", "## 最终汇报"):
        assert part in text


# ---------------------------------------------------------------- kb_search 工具

def test_kb_tool_results_scope_and_errors(company):
    company.write_file("docs/PRD.md", "# PRD\n\n## 退款\n\n订单取消后 24 小时内退款。", author="pm")
    text, summary = call_kb(company, json.dumps({"query": "退款多久到账"}))
    assert "docs/PRD.md 第" in text and "24 小时" in text and "引用时写明出处" in text
    assert summary["hits"] == 1 and summary["query"] == "退款多久到账" and summary["sources"][0].startswith("docs/PRD.md")
    assert "没有找到" in call_kb(company, json.dumps({"query": "退款多久到账", "scope": "records"}))[0]
    assert call_kb(company, json.dumps({"query": ""}))[0].startswith("错误")
    assert call_kb(company, json.dumps({"query": "x", "scope": "web"}))[0].startswith("错误")
    assert call_kb(company, "{bad")[0].startswith("错误")


def test_worker_gets_kb_tool_only_when_enabled(company):
    class LLM:
        def __init__(self):
            self.tools, self.msgs = [], []
            self.replies = [reply("先查一下", [call("k1", "kb_search", query="退款规则")]), reply("查到了")]

        def chat(self, messages, tools=None, **kw):
            self.tools.append([t["function"]["name"] for t in tools or []])
            self.msgs.append(messages)
            return self.replies.pop(0)

    company.write_file("docs/PRD.md", "# PRD\n\n退款规则：7 天内无理由退款。", author="pm")
    llm = LLM()
    company.llm_for = lambda role_id: llm
    events = []
    out = run(company, "pm", [{"role": "user", "content": "写方案"}], events.append, "m1", "all", threading.Event(), 4)
    assert "kb_search" in llm.tools[0] and "kb_search" in llm.msgs[0][0]["content"]   # 工具和提示词里的规则
    assert "7 天内无理由退款" in llm.msgs[1][-1]["content"]
    chip = next(e for e in events if e["type"] == "tool" and e["name"] == "kb_search")
    assert chip["hits"] == 1 and out["tools"][0]["query"] == "退款规则"

    company.set_kb(False)
    assert company.kb is False and "kb_search" not in system_prompt(company, "pm")
    llm.replies = [reply("好")]
    run(company, "pm", [{"role": "user", "content": "写方案"}], events.append, "m2", "all", threading.Event(), 4)
    assert "kb_search" not in llm.tools[-1]


def test_old_companies_default_to_kb_on(company):
    company.meta.pop("kb")
    assert company.kb is True and company.public_meta()["kb"] is True


# ---------------------------------------------------------------- 规划时自动召回（跨指令记忆）

def test_second_directive_recalls_first_directive(tmp_path):
    env = Env(tmp_path, roles=("pm", "engineer"), auto_approve=True,
              plans=[{"summary": "写需求", "tasks": [{"key": "t1", "title": "写 PRD", "detail": "产出 docs/PRD.md",
                                                     "assignee": "pm", "depends_on": []}]},
                     {"summary": "做退款", "tasks": [{"key": "t1", "title": "退款方案", "detail": "产出 docs/refund.md",
                                                     "assignee": "engineer", "depends_on": []}]}])
    env.start("做一个记账 App，金额一律用分存储，不做外币", did="d1")
    env.world.first_calls = {"engineer": [call("k1", "kb_search", query="金额 存储 外币", scope="records")]}
    env.start("给记账 App 加退款功能，金额和外币怎么处理", did="d2")
    plans = [m for role, m in env.world.calls if role is None and "规划助手" in m[0]["content"]]
    assert "公司知识库里和这条指令相关的记录" not in plans[0][1]["content"]   # 第一条指令：还没有记忆
    second = plans[1][1]["content"]
    assert "公司知识库里和这条指令相关的记录" in second and "金额一律用分存储，不做外币" in second
    plan_msg = [m for m in env.chans()["all"] if m.get("kind") == "plan"][-1]
    assert any("指令结论" in x for x in plan_msg["recalled"])
    assert not any("退款功能" in x for x in plan_msg["recalled"])          # 这条指令自己的结论不算
    # 执行的同事用 kb_search 查得到上一条指令的结论
    tool_result = next(m[-1]["content"] for role, m in env.world.calls if role == "engineer" and m[-1]["role"] == "tool")
    assert "指令结论" in tool_result and "金额一律用分存储" in tool_result
    env.engine.close()


def test_recall_is_off_when_kb_disabled(tmp_path):
    env = Env(tmp_path, roles=("pm",), auto_approve=True, company_kw={"kb": False})
    env.start("做一个记账 App", did="d1")
    env.start("记账 App 加报表", did="d2")
    plans = [m for role, m in env.world.calls if role is None and "规划助手" in m[0]["content"]]
    assert all("公司知识库" not in p[1]["content"] for p in plans)
    assert "kb_search" not in env.world.tools_seen[-1][1]
    env.engine.close()


# ---------------------------------------------------------------- 编码 agent 的检索工具

@pytest.mark.skipif(not coding.AVAILABLE, reason="没有安装 Coding-agent")
def test_code_agent_gets_kb_tool(company, monkeypatch):
    company.write_file("docs/PRD.md", "# PRD\n\n金额一律用分存储。", author="pm")
    tool = coding.KBSearchTool(coding.Workspace(company.files_root), company)
    assert tool.schema()["function"]["name"] == "kb_search" and not tool.needs_permission
    assert "金额一律用分存储" in tool.run(query="金额怎么存")

    seen = {}

    class FakeAgent:
        def __init__(self, llm, tools, *a, **kw):
            seen["tools"] = set(tools)
            self.steps = 0
            self.usage = {"prompt": 0, "completion": 0, "cached": 0}

        def run(self, task):
            pass

    monkeypatch.setattr(coding, "CodingAgent", FakeAgent)
    monkeypatch.setattr(company, "llm_for", lambda role_id: None)
    coding.run_code_agent(company, "pm", "实现", lambda ev: None, {}, threading.Event(), False)
    assert "kb_search" in seen["tools"] and "run_command" not in seen["tools"]
    company.set_kb(False)
    coding.run_code_agent(company, "pm", "实现", lambda ev: None, {}, threading.Event(), False)
    assert "kb_search" not in seen["tools"]


def test_kb_tool_schema_matches_scopes():
    assert KB_TOOL["function"]["parameters"]["properties"]["scope"]["enum"] == ["all", "files", "records", "library"]
    assert isinstance(KnowledgeBase, type)
